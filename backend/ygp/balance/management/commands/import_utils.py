import re
import pandas as pd
from balance.models import BalanceSheetEntry, BalanceSheetYearValue, CompanySymbol

PERCENTAGE_FIELDS = [
    "Return on Equity",
    "Return on Capital Emp",
    "OPM",
]

TARGET_KEYWORDS = {
    "Sales": "Sales",
    "OPM": "OPM",
    "Net profit": "Net profit",
    "EPS": "EPS",
    "Return on Equity": "ROE",
    "Return on Capital Emp": "ROCE",
    "Return on Capital Employed": "ROCE",
}

import datetime
import pandas as pd
import re

def normalize_column_name(col):
    # Excel date
    if isinstance(col, (pd.Timestamp, datetime.date, datetime.datetime)):
        return col.year

    # String representations
    if isinstance(col, str):
        c = col.strip()

        # Mar-18 / Mar 18
        m = re.match(r"^(Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)[\-\s]?(\d{2,4})$", c, re.I)
        if m:
            y = int(m.group(2))
            return 2000 + y if y < 100 else y

        # FY18 / FY2018
        m = re.match(r"^FY[\-\s]?(\d{2,4})$", c, re.I)
        if m:
            y = int(m.group(1))
            return 2000 + y if y < 100 else y

        # "2018-03-31"
        m = re.match(r"^(19|20)\d{2}", c)
        if m:
            return int(c[:4])

    return None

def detect_year_columns(df):
    """Return {year: column_name}"""
    detected = {}
    for col in df.columns:
        year = normalize_column_name(col)
        if year:
            detected[year] = col
    return detected

def detect_trailing_column(df):
    """Return the trailing column name if exists"""
    for c in df.columns:
        if isinstance(c, str) and ("trailing" in c.strip().lower() or "ttm" in c.strip().lower()):
            return c
    return None

from openpyxl import load_workbook
import pandas as pd

def read_excel_values(path, sheet_name, header_row):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet_name]

    data = []
    for row in ws.iter_rows(values_only=True):
        data.append(list(row))

    df = pd.DataFrame(data)
    df.columns = df.iloc[header_row]
    df = df.iloc[header_row + 1:]
    df.reset_index(drop=True, inplace=True)
    return df
import logging
import os
import re
import logging
import os
from balance.models import CompanySymbol

logger = logging.getLogger(__name__)

import os
from balance.models import CompanySymbol


def resolve_company_from_filename(fname):
    print("[RESOLVER 1] Function entered", flush=True)

    filename = os.path.splitext(os.path.basename(fname))[0].strip()
    print(f"[RESOLVER 2] Filename: {filename!r}", flush=True)

    print("[RESOLVER 3] Querying database...", flush=True)

    matches = CompanySymbol.objects.filter(name__iexact=filename)

    print("[RESOLVER 4] Query completed", flush=True)

    count = matches.count()
    print(f"[RESOLVER 5] Matches: {count}", flush=True)

    if count == 0:
        print(f"[RESOLVER ERROR] No mapping for {filename!r}", flush=True)
        return None

    if count > 1:
        raise RuntimeError(
            f"Multiple CompanySymbol mappings found for {filename!r}"
        )

    company = matches.first()
    print(
        f"[RESOLVER 6] Matched: {company.name} | {company.symbol}",
        flush=True,
    )

    return company
import pandas as pd
import re
"""
def read_excel_values(path, sheet_name, header_row):
    wb = load_workbook(path, data_only=False)  # ❗ formulas enabled
    ws = wb[sheet_name]

    data = []
    headers = []

    for r_idx, row in enumerate(ws.iter_rows()):
        row_vals = []
        for cell in row:
            if r_idx == header_row:
                # ---- HEADER CELL ----
                if cell.value is None and cell.data_type == "f":
                    headers.append(None)
                elif isinstance(cell.value, str) and cell.value.startswith("="):
                    # Formula header
                    m = re.match(r"='([^']+)'!([A-Z]+\d+)", cell.value)
                    if m:
                        ref_sheet, ref_cell = m.groups()
                        headers.append(
                            wb[ref_sheet][ref_cell].value
                        )
                    else:
                        headers.append(cell.value)
                else:
                    headers.append(cell.value)
            else:
                row_vals.append(cell.value)
        if r_idx != header_row:
            data.append(row_vals)

    df = pd.DataFrame(data)
    df.columns = headers
    df.reset_index(drop=True, inplace=True)
    return df

"""

import pandas as pd
import re

def read_excel_values(path, sheet_name, header_row):
    # Workbook 1: formulas (for headers)
    wb_f = load_workbook(path, data_only=False)
    ws_f = wb_f[sheet_name]

    # Workbook 2: values (for data)
    wb_v = load_workbook(path, data_only=True)
    ws_v = wb_v[sheet_name]

    headers = []
    data = []

    for col_idx, cell in enumerate(ws_f[header_row + 1]):
        if isinstance(cell.value, str) and cell.value.startswith("="):
            m = re.match(r"='([^']+)'!([A-Z]+\d+)", cell.value)
            if m:
                ref_sheet, ref_cell = m.groups()
                headers.append(wb_f[ref_sheet][ref_cell].value)
            else:
                headers.append(cell.value)
        else:
            headers.append(cell.value)

    for r_idx in range(header_row + 2, ws_v.max_row + 1):
        row_vals = []
        for c_idx in range(1, len(headers) + 1):
            row_vals.append(ws_v.cell(row=r_idx, column=c_idx).value)
        data.append(row_vals)

    df = pd.DataFrame(data, columns=headers)
    df.reset_index(drop=True, inplace=True)
    return df


import datetime
import re

MONTH_MAP = {
    "mar": 3,
    "jun": 6,
    "sep": 9,
    "dec": 12,
}

def detect_quarter_columns(df):
    quarter_cols = {}

    for col in df.columns:
        col_clean = col.strip().lower()

        # Match Mar-24, Mar 24, Mar_24 etc
        match = re.search(r"(mar|jun|sep|dec)[\s\-_]?(\d{2})", col_clean)
        if match:
            month_str, year_suffix = match.groups()
            month = MONTH_MAP.get(month_str)

            year = 2000 + int(year_suffix)

            quarter_date = datetime.date(year, month, 1)

            quarter_cols[quarter_date] = col

    return quarter_cols
